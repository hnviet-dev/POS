#!/usr/bin/env python3
"""Generate enhanced thesis summary v2 — deep business flow, payment & AI chatbot."""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


def set_run_font(run, name="Times New Roman", size=13, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=14, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        set_run_font(run, size=13, bold=True)
    return p


def add_body(doc, text, indent=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.left_indent = Cm(1.27)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run)
    return p


def setup_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
    return doc


def build_document():
    doc = setup_document()

    # ===== TRANG BÌA =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line, sz, bold in [
        ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ GTVT\n", 13, True),
        ("KHOA CÔNG NGHỆ THÔNG TIN\n\n", 13, True),
        ("TÓM TẮT ĐỒ ÁN TỐT NGHIỆP\n\n", 16, True),
        (
            "XÂY DỰNG HỆ THỐNG ĐIỂM BÁN HÀNG (POS)\n"
            "ĐA CHI NHÁNH CHO CHUỖI CỬA HÀNG BÁN LẺ\n"
            "(ZOSH POS)\n",
            14,
            True,
        ),
    ]:
        r = title.add_run(line)
        set_run_font(r, size=sz, bold=bold)

    doc.add_paragraph()
    for label, value in [
        ("Họ và tên sinh viên:", "..............................................................."),
        ("Mã sinh viên:", ".........................................................................."),
        ("Lớp:", "..........................."),
        ("Ngành đào tạo:", "Công nghệ thông tin"),
        ("Khoa:", "Công nghệ thông tin"),
        ("Giảng viên hướng dẫn:", "ThS. Phạm Văn A"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        r1 = p.add_run(f"{label} ")
        set_run_font(r1, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2)

    doc.add_page_break()

    # ===== 1–4: THÔNG TIN, TÊN, MỤC TIÊU, PHẠM VI =====
    add_heading(doc, "1. THÔNG TIN SINH VIÊN")
    for line in [
        "Họ và tên: ...............................................................",
        "Mã sinh viên: ..........................................................................",
        "Lớp: ...........................",
        "Ngành đào tạo: Công nghệ thông tin",
        "Khoa: Công nghệ thông tin",
        "Giảng viên hướng dẫn: ThS. Phạm Văn A",
    ]:
        add_body(doc, line, indent=True)

    add_heading(doc, "2. TÊN ĐỀ TÀI", level=2)
    add_body(
        doc,
        "Xây dựng hệ thống điểm bán hàng (POS) đa chi nhánh cho chuỗi cửa hàng bán lẻ — Zosh POS.",
        indent=True,
    )

    add_heading(doc, "3. MỤC TIÊU CỦA ĐỀ TÀI", level=2)
    add_body(
        doc,
        "Đề tài hướng tới xây dựng nền tảng quản lý bán hàng SaaS multi-tenant trên web, cho phép nhiều "
        "chuỗi cửa hàng bán lẻ đăng ký và vận hành độc lập trên cùng một hệ thống. Hệ thống giải quyết "
        "bài toán quản lý kho thủ công, thiếu đồng bộ dữ liệu giữa chi nhánh, báo cáo không theo thời gian "
        "thực và phân quyền nhân sự chưa rõ ràng. Ngoài các nghiệp vụ POS cốt lõi, đề tài tích hợp "
        "thanh toán QR thực tế qua SePay/VietQR và AI Chatbot hỗ trợ phân tích kinh doanh bằng ngôn ngữ tự nhiên.",
        indent=True,
    )

    add_heading(doc, "4. GIỚI HẠN VÀ PHẠM VI CỦA ĐỀ TÀI", level=2)
    add_body(doc, "Phạm vi:", indent=True, bold=True)
    for item in [
        "Ứng dụng web full-stack: Spring Boot 3.5.3 (backend) + React 19/Vite 7 (frontend), MySQL 8.0.",
        "7 vai trò người dùng với RBAC và lọc dữ liệu theo store_id/branch_id.",
        "Nghiệp vụ: onboarding cửa hàng, subscription SaaS, quản lý chi nhánh/sản phẩm/kho/nhân viên, "
        "bán hàng POS, ca làm việc, hoàn trả, báo cáo, thanh toán đa hình thức, AI Chatbot.",
        "Tích hợp: Cloudinary, SePay (QR), Razorpay/Stripe (subscription), Google Gemini AI, Gmail SMTP.",
    ]:
        add_bullet(doc, item)
    add_body(doc, "Giới hạn:", indent=True, bold=True)
    for item in [
        "Chưa triển khai production cloud; demo trên localhost/Docker.",
        "Thanh toán thẻ tại quầy xác nhận thủ công (không tích hợp trực tiếp máy POS vật lý).",
        "Chưa phát triển ứng dụng mobile native và module thương mại điện tử.",
    ]:
        add_bullet(doc, item)

    # ===== 5. NỘI DUNG THỰC HIỆN — PHẦN TRỌNG TÂM =====
    add_heading(doc, "5. NỘI DUNG THỰC HIỆN", level=2)

    add_body(doc, "5.1. Tổng quan kiến trúc hệ thống", bold=True)
    add_body(
        doc,
        "Zosh POS áp dụng mô hình Client–Server với kiến trúc phân tầng: Presentation (React SPA) → "
        "REST API (Spring Boot Controller) → Business Logic (Service) → Persistence (JPA Repository) → MySQL. "
        "Mô hình multi-tenant đảm bảo mỗi cửa hàng (Store) hoạt động độc lập, dữ liệu phân tách logic qua "
        "store_id và branch_id. Xác thực JWT stateless, phân quyền @PreAuthorize kết hợp lọc dữ liệu theo "
        "ngữ cảnh người dùng đăng nhập.",
        indent=True,
    )

    add_body(doc, "5.2. Sáu luồng nghiệp vụ chính của hệ thống", bold=True)

    flows = [
        (
            "Luồng 1 — Onboarding cửa hàng (SaaS)",
            "Store Admin đăng ký tài khoản (POST /auth/signup) → tạo cửa hàng (POST /api/stores, status=PENDING) "
            "→ Super Admin duyệt (PUT /api/stores/{id}/moderate, chuyển ACTIVE) → Store Admin đăng ký gói dịch vụ "
            "(POST /api/subscriptions/subscribe) → thanh toán qua Razorpay/Stripe → xác minh (POST /api/payments/verify) "
            "→ subscription chuyển ACTIVE. Luồng này chứng minh mô hình SaaS: nhiều doanh nghiệp cùng sử dụng một nền tảng.",
        ),
        (
            "Luồng 2 — Thiết lập dữ liệu nền (Store Admin)",
            "Sau khi được duyệt, Store Admin tạo chi nhánh (POST /api/branches), danh mục và sản phẩm "
            "(POST /api/categories, POST /api/products — upload ảnh Cloudinary), thêm nhân viên và phân quyền "
            "(POST /api/stores/add/employee — gán Branch Manager/Cashier vào chi nhánh). Tồn kho (inventories) "
            "được gắn theo từng chi nhánh, mỗi sản phẩm có bản ghi quantity riêng.",
        ),
        (
            "Luồng 3 — Bán hàng tại quầy (Cashier) — LUỒNG CỐT LÕI",
            "Thu ngân đăng nhập → bắt buộc mở ca (POST /api/shift-reports/start, mỗi thu ngân 1 ca/ngày) "
            "→ chọn sản phẩm trên giao diện POS → giỏ hàng tính realtime bằng Redux (subtotal, thuế VAT, giảm giá) "
            "→ thanh toán (CASH/CARD/QR) → POST /api/orders → backend validate giá từ DB, lưu order_items, "
            "trừ tồn kho → in biên lai. Điểm nổi bật: tính tiền tại frontend để không latency, backend validate "
            "chống gian lận khi checkout.",
        ),
        (
            "Luồng 4 — Tạm giữ đơn và hoàn trả",
            "Hold Order: Redux action holdOrder lưu giỏ hàng vào mảng holdOrders, phục vụ khách chưa thanh toán ngay. "
            "Refund: Cashier chọn đơn (POST /api/refunds) → tạo bản ghi refund, đổi order status=REFUNDED → "
            "số tiền hoàn được trừ vào báo cáo ca làm việc. Phản ánh nghiệp vụ thực tế tại quầy bán lẻ.",
        ),
        (
            "Luồng 5 — Kết ca và đối soát (End Shift)",
            "PATCH /api/shift-reports/end → backend query tất cả Order và Refund của thu ngân trong khoảng "
            "shift_start → shift_end → tính totalSales, totalRefunds, netSales, top 5 sản phẩm bán chạy, "
            "phân loại theo CASH/CARD/UPI → lưu shift_report. Thu ngân in phiếu bàn giao tiền mặt cuối ca.",
        ),
        (
            "Luồng 6 — Báo cáo và phân tích (Dashboard)",
            "Store Admin/Branch Manager xem dashboard với Recharts: doanh thu 7 ngày, top sản phẩm, hiệu suất "
            "thu ngân, phân bổ phương thức thanh toán. Dữ liệu từ BranchAnalyticsService với SQL aggregate "
            "theo thời gian và chi nhánh. Hỗ trợ xuất báo cáo biểu đồ PDF.",
        ),
    ]
    for name, desc in flows:
        add_body(doc, name, bold=True)
        add_body(doc, desc, indent=True)

    doc.add_page_break()

    # ===== 5.3 THANH TOÁN — PHẦN SÂU =====
    add_body(doc, "5.3. Hệ thống thanh toán — Phân tích chi tiết", bold=True)
    add_body(
        doc,
        "Hệ thống có hai lớp thanh toán: (A) Thanh toán tại quầy khi bán hàng POS — 3 phương thức; "
        "(B) Thanh toán gói dịch vụ SaaS — qua Razorpay/Stripe. Đây là điểm kỹ thuật nổi bật của đồ án.",
        indent=True,
    )

    add_body(doc, "A. Thanh toán tại quầy (POS Payment)", bold=True)

    add_body(doc, "(1) Tiền mặt (CASH)", bold=True)
    add_body(
        doc,
        "Thu ngân nhập số tiền khách đưa → hệ thống tính tiền thối (cashAmount − total) → gợi ý mệnh giá "
        "nhanh (10k–500k) → validate tiền khách đưa ≥ tổng tiền → xác nhận → POST /api/orders với "
        "paymentType=CASH → hiển thị biên lai. Không cần tích hợp bên ngoài, phù hợp cửa hàng nhỏ.",
        indent=True,
    )

    add_body(doc, "(2) Thẻ ngân hàng (CARD)", bold=True)
    add_body(
        doc,
        "Khách quẹt thẻ qua máy POS vật lý (ngoài hệ thống) → thu ngân bấm xác nhận thủ công "
        "→ POST /api/orders với paymentType=CARD. Thiết kế pragmatic: không tích hợp SDK máy quẹt thẻ "
        "vì chi phí và phức tạp triển khai, phù hợp cửa hàng đã có thiết bị riêng.",
        indent=True,
    )

    add_body(doc, "(3) QR Chuyển khoản — SePay/VietQR (UPI) — TÍCH HỢP THỰC TẾ", bold=True)
    add_body(
        doc,
        "Đây là phương thức thanh toán kỹ thuật phức tạp và nổi bật nhất. Luồng hoạt động:",
        indent=True,
    )
    qr_steps = [
        "Frontend tạo mã đơn tham chiếu: qrOrderRef = \"DH\" + 7 chữ số timestamp (VD: DH4610579).",
        "POST /api/qr-payments/register — backend lưu session vào ConcurrentHashMap (orderRef, amount, status=PENDING).",
        "Hiển thị QR từ SePay API (qr.sepay.vn/img) nhúng sẵn số tài khoản VA, ngân hàng MBBank, số tiền và nội dung chuyển khoản.",
        "Khách quét QR → chuyển khoản → MBBank ghi nhận → SePay cập nhật giao dịch.",
        "Backend @Scheduled poll SePay API mỗi 4 giây (GET /transactions/list?since_id=...) — tìm mã DH trong transaction_content bằng regex.",
        "Khớp mã + số tiền (±1000đ) → confirmPayment() → session.markPaid().",
        "Frontend poll GET /api/qr-payments/status/{ref} mỗi 3 giây → nhận PAID → tự động gọi processPayment() → POST /api/orders → DELETE session → mở biên lai.",
    ]
    for step in qr_steps:
        add_bullet(doc, step)

    add_body(
        doc,
        "Lý do chọn API Polling thay vì Webhook: Webhook yêu cầu URL public (không khả thi trên localhost); "
        "Polling cho phép backend chủ động hỏi SePay, chạy được ngay khi demo đồ án. ConcurrentHashMap "
        "đảm bảo thread-safe khi nhiều khách thanh toán QR đồng thời. Trade-off: session in-memory mất khi "
        "backend restart — production nên dùng Redis.",
        indent=True,
    )

    add_body(doc, "B. Thanh toán gói dịch vụ SaaS (Subscription Payment)", bold=True)
    add_body(
        doc,
        "Store Admin chọn gói (Starter/Pro...) → POST /api/subscriptions/subscribe → SubscriptionServiceImpl "
        "tạo subscription status=PENDING → PaymentServiceImpl.initiatePayment() tạo checkout URL qua "
        "Razorpay hoặc Stripe → frontend mở tab thanh toán → sau khi thanh toán, POST /api/payments/verify "
        "xác minh payment_id với cổng → nếu captured/success → Payment=SUCCESS, Subscription=ACTIVE. "
        "Gói dịch vụ giới hạn số chi nhánh, sản phẩm, nhân viên theo plan_id.",
        indent=True,
    )

    doc.add_page_break()

    # ===== 5.4 AI CHATBOT =====
    add_body(doc, "5.4. AI Chatbot trợ lý thông minh — Phân tích chi tiết", bold=True)
    add_body(
        doc,
        "AI Chatbot là tính năng đổi mới, tích hợp Google Gemini 2.5 Flash vào giao diện POS, cho phép "
        "người dùng hỏi bằng tiếng Việt và nhận câu trả lời dựa trên dữ liệu thực từ database — không phải "
        "câu trả lời generic. Widget floating (AiChatWidget.jsx) hiển thị cho mọi user đã đăng nhập.",
        indent=True,
    )

    add_body(doc, "Kiến trúc AI Chatbot:", bold=True)
    arch_steps = [
        "User gõ câu hỏi (VD: \"Doanh thu hôm nay bao nhiêu?\") trên AiChatWidget.",
        "POST /api/ai/chat kèm JWT → AiChatController → AiChatServiceImpl.",
        "Xác thực user từ JWT → xác định role, store_id, branch_id.",
        "Thu thập context từ MySQL theo phân quyền: tổng đơn/doanh thu, doanh thu hôm nay vs hôm qua, top 5 SP, doanh thu theo chi nhánh, SP sắp hết hàng (≤10).",
        "Xây dựng system prompt tiếng Việt + context → gọi Google GenAI SDK (gemini-2.5-flash).",
        "Trả JSON {reply, timestamp} → hiển thị trong chat bubble với markdown cơ bản.",
    ]
    for step in arch_steps:
        add_bullet(doc, step)

    add_body(doc, "Phân quyền dữ liệu context:", bold=True)
    for item in [
        "Store Admin/Manager: toàn bộ dữ liệu cửa hàng — tất cả chi nhánh, đơn hàng, tồn kho.",
        "Branch Manager/Admin: chỉ dữ liệu chi nhánh được gán.",
        "Cashier: context ca làm việc (mở rộng tương lai).",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Tính năng nổi bật:", bold=True)
    for item in [
        "Chat tự nhiên tiếng Việt — trả lời dựa trên dữ liệu POS thực, không hallucinate số liệu.",
        "Dashboard Insight (GET /api/ai/dashboard-insight): 1-click phân tích tổng quan kinh doanh tự động.",
        "8 câu hỏi gợi ý nhanh (GET /api/ai/suggestions): doanh thu, top SP, tồn kho, so sánh ngày...",
        "Bảo mật: API key Gemini chỉ ở server (application.yml), không lộ frontend; mọi request qua JWT.",
        "Fallback mode khi chưa cấu hình API key — hiển thị hướng dẫn thay vì crash.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "5.5. Phân tích sâu luồng POS và quản lý ca", bold=True)
    add_body(
        doc,
        "Luồng bán hàng POS là trọng tâm kỹ thuật của đồ án, kết hợp xử lý realtime tại frontend và "
        "validate nghiêm ngặt tại backend. Khi thu ngân click sản phẩm, Redux cartSlice thực hiện addItem — "
        "tăng quantity, tính subtotal = Σ(sellingPrice × qty), áp dụng discountAmount và taxRate, "
        "total = subtotal − discount + tax. Toàn bộ diễn ra trong vài millisecond, không gọi API.",
        indent=True,
    )
    add_body(
        doc,
        "Khi checkout, OrderServiceImpl lấy cashier và branch từ JWT, duyệt từng item: đọc sellingPrice "
        "thực tế từ bảng products (không tin giá frontend gửi lên), nhân quantity, lưu order_items, "
        "cập nhật inventories giảm quantity. Cơ chế này ngăn thu ngân hoặc attacker sửa giá qua DevTools.",
        indent=True,
    )
    add_body(
        doc,
        "Quản lý ca (Shift Report): bảng shift_report lưu shift_start, shift_end, total_sales, total_refunds, "
        "net_sales, total_orders. Khi Start Shift, kiểm tra ca đang mở (shift_end = null) — nếu có thì reject. "
        "Khi End Shift, backend aggregate Order và Refund trong khoảng thời gian ca, tính topSellingProducts "
        "và paymentSummaries (CASH/CARD/UPI). Đây là cơ chế đối soát tiền mặt cuối ngày — nghiệp vụ bắt buộc "
        "tại mọi cửa hàng bán lẻ chuyên nghiệp.",
        indent=True,
    )

    add_body(doc, "5.6. Các điểm kỹ thuật nổi bật khác", bold=True)
    highlights = [
        "Multi-tenant SaaS: nhiều cửa hàng trên một nền tảng, phân tách dữ liệu logic, Super Admin giám sát toàn hệ thống.",
        "RBAC + Data Filtering: @PreAuthorize chặn API + SQL WHERE branch_id/store_id theo user hiện tại.",
        "Giỏ hàng Redux realtime: subtotal/tax/discount tính local, backend validate giá khi checkout — cân bằng UX và bảo mật.",
        "Quản lý ca làm việc (Shift): chặn bán hàng khi chưa mở ca, chống trùng ca, chốt sổ tự động khi kết ca.",
        "Thanh toán QR SePay: polling song song FE/BE, khớp giao dịch ngân hàng tự động — demo ấn tượng trước hội đồng.",
        "AI Chatbot Gemini: phân tích kinh doanh bằng ngôn ngữ tự nhiên — điểm cộng lớn khi bảo vệ đồ án.",
        "Dashboard Recharts + xuất PDF: báo cáo trực quan theo thời gian, chi nhánh, thu ngân.",
    ]
    for h in highlights:
        add_bullet(doc, h)

    add_body(doc, "5.7. Bảo mật và phân quyền dữ liệu", bold=True)
    add_body(
        doc,
        "Hệ thống áp dụng defense in depth: (1) JWT stateless — token hết hạn 24h, JwtValidator filter mọi request; "
        "(2) BCrypt mã hóa mật khẩu; (3) @PreAuthorize trên Controller — Cashier không gọi được API tạo sản phẩm; "
        "(4) Service layer lọc dữ liệu — query Order luôn có WHERE branch_id = currentUser.branchId, "
        "Store Admin chỉ thấy store_id của mình. Chi nhánh A không thể xem đơn hàng chi nhánh B dù cố gọi API trực tiếp.",
        indent=True,
    )

    doc.add_page_break()

    # ===== 6. KẾT QUẢ =====
    add_heading(doc, "6. KẾT QUẢ THỰC HIỆN", level=2)
    add_body(
        doc,
        "Sau quá trình thực hiện, đề tài đã hoàn thành hệ thống Zosh POS với đầy đủ nghiệp vụ và các tính năng "
        "kỹ thuật nổi bật:",
        indent=True,
    )
    results = [
        "Hệ thống POS đa chi nhánh hoàn chỉnh: backend 19 REST Controllers, 13 Service, frontend React SPA đa vai trò.",
        "6 luồng nghiệp vụ chính vận hành ổn định: onboarding → setup → bán hàng → hold/refund → kết ca → báo cáo.",
        "3 phương thức thanh toán tại quầy (CASH/CARD/QR SePay) + thanh toán subscription Razorpay/Stripe.",
        "Tích hợp AI Chatbot Google Gemini 2.5 Flash với 3 API endpoints, phân quyền context theo role.",
        "Cơ sở dữ liệu MySQL 15+ bảng, ERD multi-tenant, JWT + BCrypt + RBAC.",
        "Giao diện shadcn/ui + Tailwind, dashboard Recharts, widget chat dark theme.",
        "Docker Compose triển khai nhanh; tài liệu kỹ thuật và kịch bản demo 6 luồng đầy đủ.",
    ]
    for r in results:
        add_bullet(doc, r)

    add_body(doc, "Bảng công nghệ chính:", bold=True)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, t in enumerate(["Tầng", "Công nghệ", "Vai trò"]):
        hdr[i].text = t
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=12, bold=True)
    rows = [
        ("Frontend", "React 19, Vite 7, Redux Toolkit", "SPA, giỏ hàng realtime, routing"),
        ("Frontend", "Tailwind, shadcn/ui, Recharts", "UI, biểu đồ dashboard"),
        ("Backend", "Spring Boot 3.5.3, Spring Security", "REST API, JWT, RBAC"),
        ("Backend", "JPA/Hibernate, MySQL 8.0", "ORM, cơ sở dữ liệu"),
        ("Thanh toán", "SePay/VietQR, Razorpay, Stripe", "QR tại quầy, subscription SaaS"),
        ("AI", "Google Gemini 2.5 Flash, google-genai SDK", "Chatbot phân tích kinh doanh"),
        ("Tích hợp", "Cloudinary, Gmail SMTP", "Ảnh SP, email reset password"),
    ]
    for row_data in rows:
        row = tbl.add_row().cells
        for i, text in enumerate(row_data):
            row[i].text = text
            for p in row[i].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=12)

    add_body(doc, "6.2. Gợi ý trình bày trước hội đồng", bold=True)
    add_body(
        doc,
        "Khi demo và bảo vệ, nên ưu tiên các luồng sau theo thứ tự ấn tượng: (1) AI Chatbot — gõ "
        "\"Phân tích tổng quan kinh doanh\" hoặc \"Doanh thu hôm nay?\" trước khi vào Dashboard; "
        "(2) Thanh toán QR SePay — khách quét QR, hệ thống tự động xác nhận sau vài giây; "
        "(3) Luồng POS đầy đủ: mở ca → bán hàng → hold order → hoàn trả → kết ca in báo cáo; "
        "(4) Multi-tenant: Super Admin duyệt cửa hàng mới trên tab ẩn danh. Mở DevTools tab Network "
        "khi thanh toán để chứng minh gọi API thật (POST /api/orders, POST /api/ai/chat).",
        indent=True,
    )

    doc.add_paragraph()

    # ===== 7. KẾT LUẬN =====
    add_heading(doc, "7. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", level=2)
    add_body(doc, "7.1. Kết luận", bold=True)
    add_body(
        doc,
        "Đề tài đã xây dựng thành công hệ thống POS đa chi nhánh Zosh POS trên nền tảng web hiện đại, "
        "đáp ứng nghiệp vụ chuỗi bán lẻ từ khâu đăng ký SaaS, thiết lập cửa hàng, bán hàng tại quầy "
        "đến báo cáo quản trị. Hai điểm nổi bật khi bảo vệ: (1) Thanh toán QR SePay với polling tự động "
        "khớp giao dịch ngân hàng — tích hợp thực tế, không mock; (2) AI Chatbot Gemini phân tích dữ liệu "
        "POS bằng tiếng Việt — tính năng đổi mới vượt trội so với POS truyền thống.",
        indent=True,
    )
    add_body(
        doc,
        "Qua đồ án, sinh viên rèn luyện phân tích hệ thống, thiết kế CSDL, phát triển full-stack, tích hợp "
        "cổng thanh toán và AI — kỹ năng cần thiết cho kỹ sư phần mềm hiện đại.",
        indent=True,
    )

    add_body(doc, "7.2. Hướng phát triển", bold=True)
    for item in [
        "Chuyển SePay QR sang Webhook khi deploy production — giảm độ trễ xuống ~1 giây.",
        "Tích hợp VNPay/MoMo cho thị trường Việt Nam; lưu session QR vào Redis thay in-memory.",
        "Mở rộng AI Chatbot: lưu lịch sử chat, render biểu đồ inline, context ca làm việc cho Cashier.",
        "Ứng dụng mobile React Native; module e-commerce đồng bộ tồn kho online/offline.",
        "Triển khai cloud (AWS/GCP) với CI/CD, monitoring, auto-scaling.",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run("Ngày ...... tháng ...... năm 2026\n\n")
    set_run_font(r)
    r = p.add_run("Sinh viên thực hiện\n\n")
    set_run_font(r, bold=True)
    r = p.add_run("....................................")
    set_run_font(r)

    return doc


if __name__ == "__main__":
    out = "/Users/Study/DOAN/z pos-source-code_1/TOM_TAT_DO_AN_TOT_NGHIEP_V2.docx"
    build_document().save(out)
    print(f"Created: {out}")
